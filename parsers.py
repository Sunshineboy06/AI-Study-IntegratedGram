# -*- coding: utf-8 -*-
"""IceNote - 笔记文件文本抽取(.txt / .md / .docx / .pdf)"""
import os


class ParseError(Exception):
    pass


def extract_text(path):
    """根据扩展名抽取纯文本,返回 (text, ext)。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".txt", ".md", ".markdown"):
        return _read_text_file(path), ext
    if ext == ".docx":
        return _read_docx(path), ext
    if ext == ".pdf":
        return _read_pdf(path), ext
    raise ParseError("暂不支持的文件格式:%s(支持 .txt / .md / .docx / .pdf)" % ext)


def _read_text_file(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8", "utf-8-sig", "gb18030", "big5"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_docx(path):
    try:
        import docx  # python-docx
    except ImportError:
        raise ParseError("解析 .docx 需要安装依赖:pip install python-docx")
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    text = "\n".join(parts)
    if not text.strip():
        raise ParseError("该 .docx 文件中没有可提取的文字内容")
    return text


def _read_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # 兼容旧包名
        except ImportError:
            raise ParseError("解析 .pdf 需要安装依赖:pip install pypdf")
    try:
        reader = PdfReader(path)
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as e:
        raise ParseError("PDF 解析失败:%s" % e)
    text = "\n".join(pages)
    if not text.strip():
        raise ParseError("该 PDF 中没有可提取的文字(可能是扫描件,OCR 不在支持范围内)")
    return text
